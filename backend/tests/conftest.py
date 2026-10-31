"""Pytest shared fixtures, in-memory database setup, test clients, and synthetic file generators."""

import io
import os
import uuid
import pytest

os.environ.setdefault("ENVIRONMENT", "test")
os.environ.setdefault("LLM_PROVIDER", "mock")
os.environ.setdefault("EMBEDDING_PROVIDER", "mock")

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
import fitz  # PyMuPDF
import docx  # python-docx

from app.db.session import Base, get_db
import app.db.models  # register all models
from app.db.models.users import User
from app.db.models.workspaces import Workspace
from app.main import app
from app.providers.none_storage import NoneStorageProvider
from app.providers.base import MockEmbeddingProvider, MockLLMProvider
import app.providers.base as providers_base

# SQLite in-memory test database
SQLALCHEMY_DATABASE_URL = "sqlite:///:memory:"

test_engine = create_engine(
    SQLALCHEMY_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


@pytest.fixture(scope="session", autouse=True)
def init_test_db():
    """Create all database schema tables in memory."""
    Base.metadata.create_all(bind=test_engine)
    yield
    Base.metadata.drop_all(bind=test_engine)


@pytest.fixture
def db():
    """Transactional database session fixture rolled back after each test."""
    connection = test_engine.connect()
    transaction = connection.begin()
    session = TestingSessionLocal(bind=connection, join_transaction_mode="create_savepoint")

    yield session

    session.close()
    if transaction.is_active:
        transaction.rollback()
    connection.close()


@pytest.fixture(autouse=True)
def override_session_local(db, monkeypatch):
    """Ensure background tasks or standalone sessionmaker calls bind to the current test connection."""
    import app.db.session
    import app.api.routes.documents

    connection = db.get_bind()
    session_factory = lambda: TestingSessionLocal(bind=connection, join_transaction_mode="create_savepoint")

    monkeypatch.setattr(app.db.session, "SessionLocal", session_factory)
    monkeypatch.setattr(app.api.routes.documents, "SessionLocal", session_factory)


@pytest.fixture(autouse=True)
def mock_storage(monkeypatch):
    """Provide clean in-memory storage provider for tests."""
    storage = NoneStorageProvider()
    monkeypatch.setattr(providers_base, "get_storage_provider", lambda: storage)
    return storage


@pytest.fixture(autouse=True)
def mock_embeddings(monkeypatch):
    """Provide deterministic mock embeddings provider."""
    provider = MockEmbeddingProvider()
    monkeypatch.setattr(providers_base, "get_embedding_provider", lambda: provider)
    return provider


@pytest.fixture(autouse=True)
def mock_llm(monkeypatch):
    """Provide deterministic mock LLM provider for pipeline tests."""
    provider = MockLLMProvider()
    monkeypatch.setattr(providers_base, "get_llm_provider", lambda: provider)
    return provider


@pytest.fixture
def user_a(db):
    """Test user A fixture."""
    existing = db.query(User).filter(User.firebase_uid == "firebase_user_a").first()
    if existing:
        return existing
    user = User(
        id=uuid.uuid4(),
        firebase_uid="firebase_user_a",
        email="usera@citelock.test",
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


@pytest.fixture
def user_b(db):
    """Test user B fixture."""
    existing = db.query(User).filter(User.firebase_uid == "firebase_user_b").first()
    if existing:
        return existing
    user = User(
        id=uuid.uuid4(),
        firebase_uid="firebase_user_b",
        email="userb@citelock.test",
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


@pytest.fixture
def workspace_a(db, user_a):
    """Workspace owned by user A."""
    ws = Workspace(
        id=uuid.uuid4(),
        owner_id=user_a.id,
        name="Workspace Alpha",
        description="Tenant A isolated workspace",
    )
    db.add(ws)
    db.commit()
    db.refresh(ws)
    return ws


@pytest.fixture
def workspace_b(db, user_b):
    """Workspace owned by user B."""
    ws = Workspace(
        id=uuid.uuid4(),
        owner_id=user_b.id,
        name="Workspace Beta",
        description="Tenant B isolated workspace",
    )
    db.add(ws)
    db.commit()
    db.refresh(ws)
    return ws


@pytest.fixture
def client(db, user_a):
    """TestClient overriding get_db and current_user to user_a."""
    def override_get_db():
        yield db

    from app.api import deps
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[deps.get_current_user] = lambda: user_a

    with TestClient(app) as test_client:
        yield test_client

    app.dependency_overrides.clear()


# --- File Generation Helpers ---

def create_sample_pdf(text: str = "This is a sample PDF document for testing CiteLock.", pages: int = 1) -> bytes:
    """Generate in-memory PDF with text layer using PyMuPDF."""
    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page(width=595, height=842)
        page.insert_text((50, 72), f"Page {i + 1} Content: {text}")
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_scanned_pdf() -> bytes:
    """Generate in-memory PDF that contains an image with NO text layer."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    # Valid PNG image bytes (1x1 pixel PNG)
    png_bytes = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x00\x00\x00\x00:~\x9bU\x00\x00\x00\nIDATx\x9cc`\x00\x00\x00\x02"
        b"\x00\x01H\xaf\xa4q\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    page.insert_image(fitz.Rect(50, 50, 250, 250), stream=png_bytes)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx(paragraphs: list[str] | None = None) -> bytes:
    """Generate in-memory DOCX document with headings and paragraphs."""
    doc = docx.Document()
    doc.add_heading("Section 1: Architecture", level=1)
    paras = paragraphs or ["CiteLock ensures factual accuracy in RAG through continuous claim auditing."]
    for p in paras:
        doc.add_paragraph(p)
    doc_io = io.BytesIO()
    doc.save(doc_io)
    return doc_io.getvalue()
